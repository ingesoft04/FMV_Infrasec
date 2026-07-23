import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "MANUAL-INSTALACION-Y-USO.md"
OUTPUT = ROOT / "docs" / "MANUAL-INSTALACION-Y-USO.docx"

NAVY = RGBColor(8, 32, 50)
CYAN = RGBColor(0, 150, 160)
GRAY = RGBColor(85, 100, 110)
LIGHT = "E8F3F5"


def set_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, 9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_inline(paragraph, text, base_size=11, color=None):
    pieces = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_font(run, base_size - 0.5, color=NAVY, name="Consolas")
        elif piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            set_font(run, base_size, bold=True, color=color)
        else:
            run = paragraph.add_run(piece)
            set_font(run, base_size, color=color)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.35)
    shade(cell, "071A28")
    set_cell_margins(cell, 130, 160, 130, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        run = p.add_run(line + ("\n" if index < len(lines) - 1 else ""))
        set_font(run, 8.5, color=RGBColor(230, 245, 247), name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, rows):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.7, 0.9, 3.75] if cols == 3 else ([1.9, 4.45] if cols == 2 else [6.35 / cols] * cols)
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                shade(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, row[c_idx] if c_idx < len(row) else "", 9.2, NAVY if r_idx == 0 else None)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in [
        ("Heading 1", 16, CYAN, 16, 8),
        ("Heading 2", 13, CYAN, 12, 6),
        ("Heading 3", 11.5, NAVY, 9, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    add_inline(header, "FMV InfraSec  |  Manual de instalación y uso", 8.5, GRAY)
    footer = section.footer.paragraphs[0]
    page_number(footer)

    cover = doc.add_paragraph()
    cover.paragraph_format.space_before = Pt(95)
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover.add_run("FMV InfraSec")
    set_font(run, 18, bold=True, color=CYAN)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("Manual de instalación y uso")
    set_font(run, 28, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(subtitle, "Portal comercial, CRM y plataforma tecnológica", 13, GRAY)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(42)
    add_inline(meta, "Versión 1.0  |  16 de julio de 2026", 10, GRAY)
    doc.add_page_break()

    doc.add_heading("Contenido", level=1)
    toc_items = [
        "Objetivo y arquitectura", "Requisitos y archivos importantes", "Instalación local con Docker",
        "Variables e integraciones", "Dominio y HTTPS", "Uso del portal de clientes",
        "Uso del CRM administrativo", "Imágenes y marca", "Operación y mantenimiento",
        "Respaldo, restauración y actualización", "Seguridad", "Solución de problemas",
        "Lista de publicación y soporte"
    ]
    for item in toc_items:
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, item, 10.5)
    doc.add_page_break()

    index = 4  # omite título y metadatos Markdown
    in_code = False
    code_lines = []
    table_rows = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                add_code(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line.replace("\u200b", ""))
            index += 1
            continue
        if line.startswith("|"):
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            index += 1
            if index < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[index]):
                index += 1
            while index < len(lines) and lines[index].startswith("|"):
                table_rows.append([cell.strip() for cell in lines[index].strip("|").split("|")])
                index += 1
            add_table(doc, table_rows)
            table_rows = []
            continue
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line), 10.5)
        elif line.startswith("- [ ] "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, "☐ " + line[6:], 10.5)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:], 10.5)
        elif line.strip():
            p = doc.add_paragraph()
            add_inline(p, line, 10.5)
        index += 1

    doc.core_properties.title = "Manual de instalación y uso - FMV InfraSec"
    doc.core_properties.subject = "Instalación, configuración, operación y mantenimiento"
    doc.core_properties.author = "FMV InfraSec"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
